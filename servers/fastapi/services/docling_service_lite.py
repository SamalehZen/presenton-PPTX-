import pdfplumber
from pptx import Presentation
from docx import Document


class DoclingService:
    def parse_to_markdown(self, file_path: str) -> str:
        ext = file_path.rsplit(".", 1)[-1].lower()
        if ext == "pdf":
            return self._parse_pdf(file_path)
        elif ext == "pptx":
            return self._parse_pptx(file_path)
        elif ext in ("docx", "doc"):
            return self._parse_docx(file_path)
        else:
            with open(file_path, "r", errors="ignore") as f:
                return f.read()

    @staticmethod
    def _parse_pdf(file_path: str) -> str:
        parts: list[str] = []
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(f"## Page {i}\n\n{text}")
        return "\n\n".join(parts)

    @staticmethod
    def _parse_pptx(file_path: str) -> str:
        prs = Presentation(file_path)
        parts: list[str] = []
        for i, slide in enumerate(prs.slides, 1):
            texts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        line = paragraph.text.strip()
                        if line:
                            texts.append(line)
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_text = " | ".join(
                            cell.text.strip() for cell in row.cells
                        )
                        texts.append(row_text)
            if texts:
                title = texts[0]
                body = "\n".join(f"- {t}" for t in texts[1:])
                parts.append(f"## Slide {i}: {title}\n\n{body}")
        return "\n\n".join(parts)

    @staticmethod
    def _parse_docx(file_path: str) -> str:
        doc = Document(file_path)
        parts: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if para.style and para.style.name.startswith("Heading"):
                level = para.style.name.replace("Heading ", "")
                try:
                    hashes = "#" * int(level)
                except ValueError:
                    hashes = "##"
                parts.append(f"{hashes} {text}")
            else:
                parts.append(text)
        for table in doc.tables:
            rows = []
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                rows.append(f"| {row_text} |")
            if rows:
                header_sep = "| " + " | ".join(
                    "---" for _ in table.rows[0].cells
                ) + " |"
                rows.insert(1, header_sep)
                parts.append("\n".join(rows))
        return "\n\n".join(parts)
